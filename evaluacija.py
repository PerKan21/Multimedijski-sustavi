"""
Evaluacija sustava za popisivanje studenata — ECAPA-TDNN
=========================================================
Mjeri preciznost prepoznavanja na temelju poznatog ground trutha
i opcijski sweepuje po parametrima kako bi pronašao optimalne postavke.

Ground truth može doći iz:
  [A] Stitch log datoteka (stitch/log/*.txt)  ← automatski, ako si koristio stitch.py
  [B] Ručna CSV datoteka                      ← format opisan u --help ispisu

Načini pokretanja:
    python evaluacija.py                 → interaktivni odabir
    python evaluacija.py --eval          → samo evaluacija s trenutnim postavkama
    python evaluacija.py --sweep         → sweep prop_decrease + faktor praga
    python evaluacija.py --gt moj_gt.csv → ručni ground truth

Pokretanje iz projekta:
    Mora se nalaziti u istom direktoriju kao main.py, model.py itd.

Format ručne CSV datoteke (--gt):
    Svaki red: naziv_snimke.wav,"Student A,Student B,Student C"
    Primjer:
        snimka1.wav,"Pero Perić,Ana Anić"
        snimka2.wav,"Pero Perić,Ivan Ivić,Maja Majić"
    Studenti koji nisu navedeni smatraju se odsutnim za tu snimku.

Instalacija:
    pip install openpyxl  (za Excel export)
"""

import os
import re
import sys
import csv
import argparse
import warnings
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

warnings.filterwarnings("ignore")

os.environ["TRANSFORMERS_OFFLINE"] = "1"
os.environ["HF_DATASETS_OFFLINE"]  = "1"

import numpy as np

import main as cfg
from model import ucitaj_model, ucitaj_bazu, izracunaj_pragove
from analiza import obradi_snimku


# ================================================================
# KONSTANTE
# ================================================================
DIR_STITCH_LOG  = os.path.join("stitch", "log")
DIR_STITCH_WAV  = os.path.join("stitch", "generirano")
REDOSLIJED_TXT  = os.path.join("snimke", "redoslijed.txt")
SEP             = "=" * 60
SEP_TANKI       = "-" * 60


# ================================================================
# DATAKLASE
# ================================================================
@dataclass
class MetrikeStudenta:
    ime:       str
    tp:        int = 0   # Prisutan i prepoznat
    fp:        int = 0   # Nije bio ali sustav ga je "prepoznao"
    fn:        int = 0   # Bio je ali sustav ga nije prepoznao
    tn:        int = 0   # Nije bio i sustav ga nije prepoznao

    @property
    def precision(self) -> float:
        return self.tp / (self.tp + self.fp) if (self.tp + self.fp) > 0 else 0.0

    @property
    def recall(self) -> float:
        return self.tp / (self.tp + self.fn) if (self.tp + self.fn) > 0 else 0.0

    @property
    def f1(self) -> float:
        p, r = self.precision, self.recall
        return 2 * p * r / (p + r) if (p + r) > 0 else 0.0

    @property
    def ukupno_snimki(self) -> int:
        return self.tp + self.fp + self.fn + self.tn


@dataclass
class RezultatEvaluacije:
    parametri:       dict
    metrike:         dict[str, MetrikeStudenta]
    tocnost_snimki:  float    # Postotak snimki s 100% točnim popisom
    macro_precision: float
    macro_recall:    float
    macro_f1:        float
    ukupno_snimki:   int
    savrsenih_snimki: int     # Snimke s 0 FP i 0 FN


# ================================================================
# PARSIRANJE STITCH LOG DATOTEKA
# ================================================================
def parsiraj_stitch_log(putanja_log: str) -> Optional[list[str]]:
    """
    Čita stitch log datoteku i vraća listu studenata koji su bili prisutni.
    Format loga generira stitch.py.

    Podržava oba formata:
      Stari: "Redoslijed govornika (N):"  — retci završavaju s (snimka.wav)
      Novi:  "Stvarni redoslijed u snimci (N):" — retci završavaju s [pocetak: Xs, trajanje: Ys]
             Uljezi su označeni s "ULJEZ [kat]" i preskaču se.

    Vraća None ako datoteka nije čitljiva ili nema studenata.
    """
    try:
        with open(putanja_log, encoding="utf-8") as f:
            sadrzaj = f.read()
    except OSError:
        return None

    # Podrži oba zaglavlja sekcije
    match = re.search(
        r"(?:Redoslijed govornika|Stvarni redoslijed u snimci)\s*\(\d+\):\n(.*?)(?:\n\n|\nIzostavljeni|\Z)",
        sadrzaj, re.DOTALL
    )
    if not match:
        return None

    prisutni = []
    for linija in match.group(1).splitlines():
        # Preskoči uljez linije (novi format označava ih s "ULJEZ")
        if "ULJEZ" in linija:
            continue
        # Stari format: "   1. Ime  (snimka.wav)"
        # Novi format:  "   1. Ime  (snimka.wav)  [pocetak: Xs, trajanje: Ys]"
        # Regex bez $ jer novi format ima tekst iza (...)
        m = re.match(r"\s*\d+\.\s+(.+?)\s+\(.*?\)", linija)
        if m:
            ime = m.group(1).strip().replace("_", " ")
            prisutni.append(ime)

    return prisutni if prisutni else None



def ucitaj_gt_iz_redoslijeda(putanja: str) -> dict[str, list[str]]:
    """
    Čita ground truth iz redoslijed.txt datoteke.
    Format: N.snimka = student1, student2, student3

    Naziv snimke se automatski traži u DIR_SNIMKE folderu.
    Tolerantan na tipfelere (višestruki zarezi, točke na kraju).
    """
    if not os.path.exists(putanja):
        return {}

    gt = {}
    with open(putanja, encoding="utf-8") as f:
        for linija in f:
            linija = linija.strip()
            if not linija or linija.startswith("#"):
                continue
            if "=" not in linija:
                continue

            naziv_kljuc, studenti_str = linija.split("=", 1)
            naziv_kljuc = naziv_kljuc.strip()  # npr. "1.snimka"

            # Čisti studente — ukloni prazne, točke, višestruke zareze
            studenti = []
            for s in re.split(r"[,]+", studenti_str):
                s = s.strip().rstrip(".").strip()
                if s:
                    studenti.append(s)

            # Pronađi odgovarajuću audio datoteku u snimke/ folderu
            putanja_snimke = None
            for f_naziv in sorted(os.listdir(cfg.DIR_SNIMKE)):
                ime_bez_ext = os.path.splitext(f_naziv)[0]
                # Podudara se s "1.snimka", "1snimka", ili direktno ime
                if (ime_bez_ext == naziv_kljuc
                        or ime_bez_ext.replace(".", "") == naziv_kljuc.replace(".", "")
                        or f_naziv == naziv_kljuc):
                    putanja_snimke = os.path.join(cfg.DIR_SNIMKE, f_naziv)
                    break

            if putanja_snimke:
                gt[putanja_snimke] = studenti
            else:
                print(f"  UPOZORENJE: Snimka za '{naziv_kljuc}' nije pronađena u '{cfg.DIR_SNIMKE}'.")

    return gt


def ucitaj_gt_iz_stitch_logova(dir_log: str, dir_wav: str) -> dict[str, list[str]]:
    """
    Automatski pronalazi parove (WAV, log) u stitch direktorijima i
    gradi ground truth dict: {naziv_wav_datoteke: [student1, student2, ...]}.

    Vraća samo parove gdje postoji i WAV i log datoteka.
    """
    if not os.path.isdir(dir_log):
        return {}

    gt = {}
    for log_dat in sorted(os.listdir(dir_log)):
        if not log_dat.endswith(".txt"):
            continue

        wav_dat  = log_dat.replace(".txt", ".wav")
        wav_put  = os.path.join(dir_wav, wav_dat)
        log_put  = os.path.join(dir_log, log_dat)

        if not os.path.exists(wav_put):
            continue

        prisutni = parsiraj_stitch_log(log_put)
        if prisutni is not None:
            gt[wav_put] = prisutni

    return gt


def ucitaj_gt_iz_csv(putanja_csv: str) -> dict[str, list[str]]:
    """
    Čita ručni ground truth iz CSV datoteke.
    Format: naziv_snimke.wav,"Student A,Student B"

    Putanja snimke je relativna na trenutni radni direktorij ili apsolutna.
    """
    gt = {}
    try:
        with open(putanja_csv, encoding="utf-8", newline="") as f:
            reader = csv.reader(f)
            for i, red in enumerate(reader):
                if not red or red[0].startswith("#"):
                    continue
                if len(red) < 2:
                    print(f"  UPOZORENJE: Redak {i+1} u CSV-u preskočen (nema stupca studenata).")
                    continue
                naziv = red[0].strip()
                studenti_str = red[1].strip()
                studenti = [s.strip() for s in studenti_str.split(",") if s.strip()]
                if not os.path.isabs(naziv):
                    # Pokušaj relativno na csv ili cwd
                    for baza_dir in [os.path.dirname(putanja_csv), os.getcwd(),
                                     cfg.DIR_SNIMKE]:
                        cand = os.path.join(baza_dir, naziv)
                        if os.path.exists(cand):
                            naziv = cand
                            break
                gt[naziv] = studenti
    except OSError as e:
        print(f"GREŠKA: Ne mogu čitati CSV: {e}")

    return gt


# ================================================================
# NORMALIZACIJA IMENA — za usporedbu baza ↔ stitch log
# ================================================================
def _normaliziraj_ime(ime: str) -> str:
    """Lowercase, podvlake → razmaci, višestruki razmaci → jedan."""
    return re.sub(r"\s+", " ", ime.replace("_", " ")).strip().lower()


def _podudaranje_studenata(prepoznati: list[str], gt_prisutni: list[str],
                            svi_studenti: list[str]) -> tuple[set, set, set, set]:
    """
    Uspoređuje prepoznate studente s ground truthom.
    Vraća (tp_set, fp_set, fn_set, tn_set) kao skupove imena studenata.
    """
    norm_prz = {_normaliziraj_ime(s): s for s in prepoznati}
    norm_gt  = {_normaliziraj_ime(s) for s in gt_prisutni}
    norm_svi = {_normaliziraj_ime(s): s for s in svi_studenti}

    tp, fp, fn, tn = set(), set(), set(), set()
    for norm_ime, orig_ime in norm_svi.items():
        je_prz   = norm_ime in norm_prz
        je_gt    = norm_ime in norm_gt
        if   je_prz and je_gt:      tp.add(orig_ime)
        elif je_prz and not je_gt:  fp.add(orig_ime)
        elif not je_prz and je_gt:  fn.add(orig_ime)
        else:                       tn.add(orig_ime)

    return tp, fp, fn, tn


# ================================================================
# POKRETANJE JEDNE EVALUACIJE
# ================================================================
def pokreni_evaluaciju(gt: dict[str, list[str]],
                       baza: dict,
                       prag_donji: float,
                       prag_gornji: float,
                       parametri: dict,
                       tiho: bool = False) -> RezultatEvaluacije:
    """
    Pokreće evaluaciju na svim GT snimkama s danim pragovima i parametrima.

    gt       = {putanja_wav: [prisutni_studenti]}
    baza     = ucitaj_bazu(...)
    tiho     = ne ispisuj detalje po snimki
    """
    svi_studenti = sorted(baza.keys())
    metrike      = {ime: MetrikeStudenta(ime=ime) for ime in svi_studenti}

    ukupno_snimki    = 0
    savrsenih_snimki = 0

    for wav_put, gt_prisutni in gt.items():
        if not os.path.exists(wav_put):
            if not tiho:
                print(f"  UPOZORENJE: Snimka ne postoji, preskačem: {wav_put}")
            continue

        naziv = os.path.basename(wav_put)

        prepoznati, segmenti, uljezi, n_gov = obradi_snimku(
            wav_put, baza, prag_donji, prag_gornji,
            sr=parametri.get("sr",              cfg.SR),
            trajanje=parametri.get("trajanje",  cfg.SEGMENT_TRAJANJE),
            preklapanje=parametri.get("prekl",  cfg.SEGMENT_PREKLAPANJE),
            prop_decrease=parametri.get("prop", cfg.SUM_PROP_DECREASE),
            vad_top_db=parametri.get("vad_db",  cfg.VAD_TOP_DB),
            vad_min_duljina=parametri.get("vad_min", cfg.VAD_MIN_DULJINA),
            vad_spajanje=parametri.get("vad_sp",    cfg.VAD_SPAJANJE),
        )

        tp_s, fp_s, fn_s, tn_s = _podudaranje_studenata(
            prepoznati, gt_prisutni, svi_studenti
        )

        for ime in tp_s: metrike[ime].tp += 1
        for ime in fp_s: metrike[ime].fp += 1
        for ime in fn_s: metrike[ime].fn += 1
        for ime in tn_s: metrike[ime].tn += 1

        savrsena = not fp_s and not fn_s
        ukupno_snimki    += 1
        savrsenih_snimki += int(savrsena)

        if not tiho:
            oznaka = "✓" if savrsena else "✗"
            print(f"  [{oznaka}] {naziv}")
            if fp_s: print(f"       FP (lažna prisutnost): {', '.join(sorted(fp_s))}")
            if fn_s: print(f"       FN (propušteni):       {', '.join(sorted(fn_s))}")

    # Makro prosjeci
    prec_list = [m.precision for m in metrike.values()]
    rec_list  = [m.recall    for m in metrike.values()]
    f1_list   = [m.f1        for m in metrike.values()]

    mac_p = np.mean(prec_list) if prec_list else 0.0
    mac_r = np.mean(rec_list)  if rec_list  else 0.0
    mac_f = np.mean(f1_list)   if f1_list   else 0.0
    toc   = savrsenih_snimki / ukupno_snimki if ukupno_snimki > 0 else 0.0

    return RezultatEvaluacije(
        parametri        = parametri,
        metrike          = metrike,
        tocnost_snimki   = toc,
        macro_precision  = mac_p,
        macro_recall     = mac_r,
        macro_f1         = mac_f,
        ukupno_snimki    = ukupno_snimki,
        savrsenih_snimki = savrsenih_snimki,
    )


# ================================================================
# ISPIS REZULTATA
# ================================================================
def ispisi_rezultat(rez: RezultatEvaluacije):
    print()
    print(SEP)
    print("  REZULTATI EVALUACIJE")
    print(SEP)

    # Tablica po studentu
    print(f"\n  {'Student':<22} {'P':>6} {'R':>6} {'F1':>6}  {'TP':>4} {'FP':>4} {'FN':>4} {'TN':>4}")
    print(f"  {SEP_TANKI}")
    for ime, m in sorted(rez.metrike.items()):
        bar_tp = "█" * m.tp
        bar_fn = "░" * m.fn
        print(
            f"  {ime:<22} {m.precision:>5.1%} {m.recall:>5.1%} {m.f1:>5.1%}"
            f"  {m.tp:>4} {m.fp:>4} {m.fn:>4} {m.tn:>4}  {bar_tp}{bar_fn}"
        )

    print(f"  {SEP_TANKI}")
    print(
        f"  {'MAKRO PROSJEK':<22} "
        f"{rez.macro_precision:>5.1%} {rez.macro_recall:>5.1%} {rez.macro_f1:>5.1%}"
    )
    print()
    print(f"  Savršene snimke:  {rez.savrsenih_snimki}/{rez.ukupno_snimki}"
          f"  ({rez.tocnost_snimki:.1%})")
    print()

    # Parametri
    p = rez.parametri
    print(f"  Parametri: prop_decrease={p.get('prop', cfg.SUM_PROP_DECREASE):.2f}  "
          f"vad_top_db={p.get('vad_db', cfg.VAD_TOP_DB)}  "
          f"prag_donji={p.get('prag_d', '?'):.4f}  "
          f"prag_gornji={p.get('prag_g', '?'):.4f}")
    print(SEP)


# ================================================================
# SWEEP — sweep po prop_decrease i faktoru praga
# ================================================================
def pokreni_sweep(gt: dict[str, list[str]], baza: dict,
                  prop_vrijednosti: list[float],
                  faktor_vrijednosti: list[float],
                  trajanje_vrijednosti: list[float] = None,
                  preklapanje_vrijednosti: list[float] = None,
                  vad_db_vrijednosti: list[float] = None,
                  vad_min_vrijednosti: list[float] = None,
                  vad_sp_vrijednosti: list[float] = None):
    """
    Full sweep po svim parametrima sustava:
      - prop_decrease (redukcija šuma)
      - faktor_gornjeg_praga
      - segment_trajanje
      - segment_preklapanje
      - vad_top_db
      - vad_min_duljina
      - vad_spajanje

    Parametri koji nisu zadani koriste defaulte iz main.py.
    """
    # Defaulti za parametre koji nisu zadani
    if trajanje_vrijednosti    is None: trajanje_vrijednosti    = [cfg.SEGMENT_TRAJANJE]
    if preklapanje_vrijednosti is None: preklapanje_vrijednosti = [cfg.SEGMENT_PREKLAPANJE]
    if vad_db_vrijednosti      is None: vad_db_vrijednosti      = [cfg.VAD_TOP_DB]
    if vad_min_vrijednosti     is None: vad_min_vrijednosti     = [cfg.VAD_MIN_DULJINA]
    if vad_sp_vrijednosti      is None: vad_sp_vrijednosti      = [cfg.VAD_SPAJANJE]

    import itertools
    kombinacije = list(itertools.product(
        prop_vrijednosti,
        faktor_vrijednosti,
        trajanje_vrijednosti,
        preklapanje_vrijednosti,
        vad_db_vrijednosti,
        vad_min_vrijednosti,
        vad_sp_vrijednosti,
    ))

    print()
    print(SEP)
    print("  FULL SWEEP PARAMETARA")
    print(SEP)
    print(f"  prop_decrease:        {prop_vrijednosti}")
    print(f"  faktor_praga:         {faktor_vrijednosti}")
    print(f"  segment_trajanje:     {trajanje_vrijednosti}")
    print(f"  segment_preklapanje:  {preklapanje_vrijednosti}")
    print(f"  vad_top_db:           {vad_db_vrijednosti}")
    print(f"  vad_min_duljina:      {vad_min_vrijednosti}")
    print(f"  vad_spajanje:         {vad_sp_vrijednosti}")
    print(f"  GT snimki:            {len(gt)}")
    print(f"  Ukupno kombinacija:   {len(kombinacije)}")
    print()

    rezultati = []
    ukupno = len(kombinacije)

    for br, (prop, faktor, traj, prekl, vad_db, vad_min, vad_sp) in enumerate(kombinacije, 1):
        print(
            f"  [{br:3d}/{ukupno}] "
            f"prop={prop:.2f} fakt={faktor:.1f} "
            f"traj={traj:.1f} prekl={prekl:.1f} "
            f"db={vad_db} min={vad_min:.2f} sp={vad_sp:.2f} ...",
            end="", flush=True
        )

        pd, pg = izracunaj_pragove(baza, faktor=faktor)
        params = {
            "prop":    prop,
            "faktor":  faktor,
            "prag_d":  pd,
            "prag_g":  pg,
            "trajanje": traj,
            "prekl":   prekl,
            "vad_db":  vad_db,
            "vad_min": vad_min,
            "vad_sp":  vad_sp,
        }

        rez = pokreni_evaluaciju(gt, baza, pd, pg, params, tiho=True)
        rezultati.append(rez)

        print(
            f"  P={rez.macro_precision:.1%}  "
            f"R={rez.macro_recall:.1%}  "
            f"F1={rez.macro_f1:.1%}  "
            f"savršene={rez.savrsenih_snimki}/{rez.ukupno_snimki}"
        )

    # Sortiraj po F1, pa po savršenim snimkama kao tiebreaker
    rezultati.sort(key=lambda r: (r.macro_f1, r.savrsenih_snimki), reverse=True)

    print()
    print(SEP)
    print("  RANG-LISTA TOP 20 (sortirano po F1)")
    print(SEP)
    print(f"  {'#':>3}  {'prop':>5}  {'fakt':>5}  {'traj':>5}  {'pre':>4}  "
          f"{'db':>4}  {'min':>5}  {'sp':>5}  {'P':>6}  {'R':>6}  {'F1':>6}  {'Savršene':>10}")
    print(f"  {SEP_TANKI}")

    for i, rez in enumerate(rezultati[:20], 1):
        p = rez.parametri
        print(
            f"  {i:>3}.  {p['prop']:>5.2f}  {p['faktor']:>5.1f}  "
            f"{p['trajanje']:>5.1f}  {p['prekl']:>4.1f}  "
            f"{p['vad_db']:>4}  {p['vad_min']:>5.2f}  {p['vad_sp']:>5.2f}  "
            f"{rez.macro_precision:>5.1%}  {rez.macro_recall:>5.1%}  "
            f"{rez.macro_f1:>5.1%}  "
            f"{rez.savrsenih_snimki:>4}/{rez.ukupno_snimki:<4}"
        )

    print(SEP)
    print()
    best = rezultati[0]
    bp   = best.parametri
    print("  ★  Optimalne postavke (po F1):")
    print(f"     SUM_PROP_DECREASE    = {bp['prop']:.2f}")
    print(f"     FAKTOR_GORNJEG_PRAGA = {bp['faktor']:.1f}")
    print(f"     SEGMENT_TRAJANJE     = {bp['trajanje']:.1f}")
    print(f"     SEGMENT_PREKLAPANJE  = {bp['prekl']:.1f}")
    print(f"     VAD_TOP_DB           = {bp['vad_db']}")
    print(f"     VAD_MIN_DULJINA      = {bp['vad_min']:.2f}")
    print(f"     VAD_SPAJANJE         = {bp['vad_sp']:.2f}")
    print(f"     → donji prag: {bp['prag_d']:.4f},  gornji prag: {bp['prag_g']:.4f}")
    print()

    return rezultati


# ================================================================
# EXCEL EXPORT
# ================================================================
def spremi_sweep_excel(rezultati: list[RezultatEvaluacije], putanja: str):
    """Sprema sweep rezultate u Excel tablicu."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    except ImportError:
        print("  UPOZORENJE: openpyxl nije instaliran, Excel export preskočen.")
        print("  Instalacija: pip install openpyxl")
        return

    wb  = Workbook()
    ws  = wb.active
    ws.title = "Sweep rezultati"
    ws.sheet_view.showGridLines = False

    tamno    = PatternFill("solid", fgColor="1e1e1e")
    srednje  = PatternFill("solid", fgColor="2a2a2a")
    zeleno   = PatternFill("solid", fgColor="2a5c2a")
    crveno   = PatternFill("solid", fgColor="5c2a2a")
    narancasto = PatternFill("solid", fgColor="5c4a1a")
    bijeli_b = Font(color="e8e8e8", bold=True, name="Calibri", size=11)
    bijeli   = Font(color="e8e8e8",             name="Calibri", size=10)
    zeleni   = Font(color="7ecf85", bold=True, name="Calibri", size=11)
    crveni   = Font(color="cf7e7e",             name="Calibri", size=10)
    nar_font = Font(color="e8943a",             name="Calibri", size=10)
    cen      = Alignment(horizontal="center", vertical="center")
    tanki    = Side(style="thin", color="444444")
    rub      = Border(left=tanki, right=tanki, top=tanki, bottom=tanki)

    # Naslov
    ws.merge_cells("A1:I1")
    c = ws.cell(1, 1, value="Evaluacija — sweep parametara")
    c.fill = tamno; c.font = Font(color="e8e8e8", bold=True, name="Calibri", size=14)
    c.alignment = cen
    ws.row_dimensions[1].height = 32

    # Zaglavlje
    zaglavlja = ["#", "prop", "faktor", "trajanje", "prekl", "vad_db", "vad_min", "vad_sp",
                 "prag_d", "prag_g", "Precision", "Recall", "F1", "Savršene"]
    sirine    = [5, 8, 7, 8, 7, 7, 7, 7, 10, 10, 10, 10, 10, 12]
    from openpyxl.utils import get_column_letter
    for col, (zag, sir) in enumerate(zip(zaglavlja, sirine), 1):
        c = ws.cell(2, col, value=zag)
        c.fill = srednje; c.font = bijeli_b; c.alignment = cen; c.border = rub
        ws.column_dimensions[get_column_letter(col)].width = sir
    ws.row_dimensions[2].height = 22

    # Podaci
    for row, rez in enumerate(rezultati, 3):
        p     = rez.parametri
        rang  = row - 2
        f1    = rez.macro_f1
        fill  = zeleno if f1 >= 0.90 else (narancasto if f1 >= 0.75 else crveno)
        font  = zeleni if f1 >= 0.90 else (nar_font  if f1 >= 0.75 else crveni)

        vrijednosti = [
            rang,
            f"{p.get('prop', 0):.2f}",
            f"{p.get('faktor', 0):.1f}",
            f"{p.get('trajanje', cfg.SEGMENT_TRAJANJE):.1f}",
            f"{p.get('prekl', cfg.SEGMENT_PREKLAPANJE):.1f}",
            f"{p.get('vad_db', cfg.VAD_TOP_DB)}",
            f"{p.get('vad_min', cfg.VAD_MIN_DULJINA):.2f}",
            f"{p.get('vad_sp', cfg.VAD_SPAJANJE):.2f}",
            f"{p.get('prag_d', 0):.4f}",
            f"{p.get('prag_g', 0):.4f}",
            f"{rez.macro_precision:.1%}",
            f"{rez.macro_recall:.1%}",
            f"{f1:.1%}",
            f"{rez.savrsenih_snimki}/{rez.ukupno_snimki}",
        ]
        for col, v in enumerate(vrijednosti, 1):
            c = ws.cell(row, col, value=v)
            c.fill = fill; c.font = font; c.alignment = cen; c.border = rub
        ws.row_dimensions[row].height = 18

    wb.save(putanja)
    print(f"  Sweep tablica spremljena: {putanja}")


def spremi_eval_excel(rez: RezultatEvaluacije, putanja: str):
    """Sprema detaljne metrike jedne evaluacije u Excel."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  UPOZORENJE: openpyxl nije instaliran.")
        return

    wb  = Workbook()
    ws  = wb.active
    ws.title = "Metrike studenata"
    ws.sheet_view.showGridLines = False

    tamno   = PatternFill("solid", fgColor="1e1e1e")
    srednje = PatternFill("solid", fgColor="2a2a2a")
    sub     = PatternFill("solid", fgColor="333333")
    zeleno  = PatternFill("solid", fgColor="2a5c2a")
    naranj  = PatternFill("solid", fgColor="5c4a1a")
    crveno  = PatternFill("solid", fgColor="5c2a2a")
    bijeli_b = Font(color="e8e8e8", bold=True, name="Calibri", size=11)
    bijeli   = Font(color="e8e8e8",             name="Calibri", size=10)
    zeleni   = Font(color="7ecf85", bold=True, name="Calibri", size=11)
    nar_f    = Font(color="e8943a",             name="Calibri", size=10)
    crveni   = Font(color="cf7e7e",             name="Calibri", size=10)
    cen      = Alignment(horizontal="center", vertical="center")
    lij      = Alignment(horizontal="left",   vertical="center")
    tanki    = Side(style="thin", color="444444")
    rub      = Border(left=tanki, right=tanki, top=tanki, bottom=tanki)

    ws.merge_cells("A1:H1")
    c = ws.cell(1, 1, value="Evaluacija sustava — metrike po studentu")
    c.fill = tamno; c.font = Font(color="e8e8e8", bold=True, name="Calibri", size=14)
    c.alignment = cen; ws.row_dimensions[1].height = 32

    p = rez.parametri
    ws.merge_cells("A2:H2")
    ts = datetime.now().strftime("%d.%m.%Y. %H:%M:%S")
    info = (f"prop={p.get('prop', cfg.SUM_PROP_DECREASE):.2f}  "
            f"faktor={p.get('faktor', cfg.FAKTOR_GORNJEG_PRAGA):.1f}  "
            f"prag_d={p.get('prag_d', 0):.4f}  prag_g={p.get('prag_g', 0):.4f}  |  {ts}")
    c = ws.cell(2, 1, value=info)
    c.fill = tamno; c.font = Font(color="888888", name="Calibri", size=9)
    c.alignment = cen; ws.row_dimensions[2].height = 18

    zaglavlja = ["Student", "Precision", "Recall", "F1", "TP", "FP", "FN", "TN"]
    sirine    = [22, 11, 11, 11, 6, 6, 6, 6]
    for col, (zag, sir) in enumerate(zip(zaglavlja, sirine), 1):
        c = ws.cell(3, col, value=zag)
        c.fill = srednje; c.font = bijeli_b; c.alignment = cen; c.border = rub
        ws.column_dimensions[get_column_letter(col)].width = sir
    ws.row_dimensions[3].height = 22

    for row, (ime, m) in enumerate(sorted(rez.metrike.items()), 4):
        fill = zeleno if m.f1 >= 0.9 else (naranj if m.f1 >= 0.7 else crveno)
        font = zeleni if m.f1 >= 0.9 else (nar_f  if m.f1 >= 0.7 else crveni)
        for col, v in enumerate([ime, f"{m.precision:.1%}", f"{m.recall:.1%}",
                                  f"{m.f1:.1%}", m.tp, m.fp, m.fn, m.tn], 1):
            c = ws.cell(row, col, value=v)
            c.fill = fill; c.font = font
            c.alignment = lij if col == 1 else cen; c.border = rub
        ws.row_dimensions[row].height = 18

    # Red makro prosjeka
    row_uk = len(rez.metrike) + 4
    for col, v in enumerate(["MAKRO PROSJEK",
                              f"{rez.macro_precision:.1%}",
                              f"{rez.macro_recall:.1%}",
                              f"{rez.macro_f1:.1%}", "", "", "", ""], 1):
        c = ws.cell(row_uk, col, value=v)
        c.fill = sub; c.font = bijeli_b
        c.alignment = lij if col == 1 else cen; c.border = rub
    ws.row_dimensions[row_uk].height = 22

    ws.freeze_panes = "A4"
    wb.save(putanja)
    print(f"  Detalji evaluacije spremljeni: {putanja}")


# ================================================================
# INTERAKTIVNI ODABIR GROUND TRUTHA
# ================================================================
def odaberi_gt(args) -> dict[str, list[str]]:
    """
    Vraća ground truth dict. Redoslijed prioriteta:
    1. --gt argument (ručni CSV)
    2. Stitch logovi (automatski, ako postoje)
    3. Korisnik ručno unosi putanje
    """
    # 1. Ručni CSV
    if hasattr(args, "gt") and args.gt:
        print(f"  Učitavam ručni GT: {args.gt}")
        gt = ucitaj_gt_iz_csv(args.gt)
        if gt:
            return gt
        print("  GREŠKA: CSV nije valjan ili je prazan.")

    # 2. Stitch logovi
    # Kombiniraj stitch logove i redoslijed.txt
    gt_kombinirani = {}

    gt_stitch = ucitaj_gt_iz_stitch_logova(DIR_STITCH_LOG, DIR_STITCH_WAV)
    if gt_stitch:
        print(f"  Pronađeno {len(gt_stitch)} GT parova iz stitch logova.")
        gt_kombinirani.update(gt_stitch)

    gt_redoslijed = ucitaj_gt_iz_redoslijeda(REDOSLIJED_TXT)
    if gt_redoslijed:
        print(f"  Pronađeno {len(gt_redoslijed)} GT parova iz redoslijed.txt.")
        gt_kombinirani.update(gt_redoslijed)

    if gt_kombinirani:
        return gt_kombinirani

    # 3. Ručni unos
    print()
    print("  Nije pronađen automatski ground truth.")
    print("  Opcije:")
    print("    [1] Unesi putanju ručnog CSV-a")
    print("    [2] Koristi snimke iz 'snimke/' s ručnim unosom prisutnih")
    odabir = input("  Odabir (1/2): ").strip()

    if odabir == "1":
        putanja = input("  Putanja CSV datoteke: ").strip()
        gt = ucitaj_gt_iz_csv(putanja)
        if gt:
            return gt
        print("  GREŠKA: Ne mogu učitati CSV.")
        return {}

    # Opcija 2 — interaktivni unos po snimci
    snimke = sorted([
        f for f in os.listdir(cfg.DIR_SNIMKE)
        if f.lower().endswith(cfg.PODRZANI_FORMATI)
        and not f.endswith("_konv.wav")
    ])
    if not snimke:
        print(f"  Nema snimki u '{cfg.DIR_SNIMKE}'.")
        return {}

    gt = {}
    print(f"\n  Za svaku snimku unesi studente koji su bili prisutni.")
    print(f"  (razdvoji zarezom, Enter = nitko nije bio prisutan)\n")
    for naziv in snimke:
        unos = input(f"  {naziv}: ").strip()
        studenti = [s.strip() for s in unos.split(",") if s.strip()]
        gt[os.path.join(cfg.DIR_SNIMKE, naziv)] = studenti

    return gt


# ================================================================
# GLAVNI PROGRAM
# ================================================================
def main():
    parser = argparse.ArgumentParser(
        description="Evaluacija sustava za popisivanje studenata.",
        formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument("--eval",  action="store_true",
                        help="Samo evaluacija s trenutnim postavkama iz main.py")
    parser.add_argument("--sweep", action="store_true",
                        help="Sweep po prop_decrease i faktoru praga")
    parser.add_argument("--gt",    type=str, default=None,
                        help="Putanja ručnog ground truth CSV-a")
    parser.add_argument("--xlsx",  action="store_true",
                        help="Spremi rezultate u Excel tablicu")
    parser.add_argument(
        "--prop-values", type=str, default="0.0,0.4,0.75",
        help="Vrijednosti prop_decrease za sweep"
    )
    parser.add_argument(
        "--faktor-values", type=str, default="1.5,2.0,2.5,3.0",
        help="Vrijednosti faktora praga za sweep"
    )
    parser.add_argument(
        "--trajanje-values", type=str, default=None,
        help="Vrijednosti segment_trajanje za sweep (npr. '1.5,2.0,3.0')"
    )
    parser.add_argument(
        "--prekl-values", type=str, default=None,
        help="Vrijednosti segment_preklapanje za sweep (npr. '0.5,1.0')"
    )
    parser.add_argument(
        "--vad-db-values", type=str, default=None,
        help="Vrijednosti vad_top_db za sweep (npr. '20,25,30')"
    )
    parser.add_argument(
        "--vad-min-values", type=str, default=None,
        help="Vrijednosti vad_min_duljina za sweep (npr. '0.3,0.5')"
    )
    parser.add_argument(
        "--vad-sp-values", type=str, default=None,
        help="Vrijednosti vad_spajanje za sweep (npr. '0.3,0.5')"
    )
    args = parser.parse_args()

    # Ako nema argumenata → interaktivni mod
    interaktivni = not (args.eval or args.sweep)

    print(SEP)
    print("  Evaluacija sustava za popisivanje studenata")
    print(SEP)

    # --- Odabir moda ---
    if interaktivni:
        print()
        print("  Odaberi mod:")
        print("  [1] Evaluacija s trenutnim postavkama (iz main.py)")
        print("  [2] Sweep parametara (prop_decrease × faktor_praga)")
        mod = input("  Odabir (1/2, Enter = 1): ").strip() or "1"
        args.eval  = (mod == "1")
        args.sweep = (mod == "2")

        if not args.xlsx:
            args.xlsx = input("  Spremi rezultate u Excel? (d/N): ").strip().lower() == "d"

    # --- Ground truth ---
    print()
    print("  Učitavanje ground trutha...")
    gt = odaberi_gt(args)
    if not gt:
        print("  GREŠKA: Nema ground truth podataka. Izlaz.")
        return

    print(f"  GT snimki: {len(gt)}")
    for put, prisutni in gt.items():
        print(f"    {os.path.basename(put)}: {', '.join(prisutni) if prisutni else '(nitko)'}")

    # --- Model i baza ---
    print()
    print(SEP)
    print("  Učitavanje modela i baze...")
    print(SEP)
    ucitaj_model()
    baza = ucitaj_bazu(
        cfg.DIR_BAZA,
        sr=cfg.SR,
        trajanje=cfg.SEGMENT_TRAJANJE,
        preklapanje=cfg.SEGMENT_PREKLAPANJE,
        prop_decrease=cfg.SUM_PROP_DECREASE,
        vad_top_db=cfg.VAD_TOP_DB,
        vad_min_duljina=cfg.VAD_MIN_DULJINA,
        vad_spajanje=cfg.VAD_SPAJANJE,
        podrzani_formati=cfg.PODRZANI_FORMATI,
        cache_putanja=cfg.CACHE_PUTANJA,
    )
    print(f"  Studenata u bazi: {len(baza)}")

    os.makedirs("rezultati", exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

    # ============================================================
    # MOD 1: EVALUACIJA
    # ============================================================
    if args.eval:
        print()
        print(SEP)
        print("  EVALUACIJA — trenutne postavke")
        print(SEP)

        prag_d, prag_g = izracunaj_pragove(baza, faktor=cfg.FAKTOR_GORNJEG_PRAGA)

        params = {
            "prop":    cfg.SUM_PROP_DECREASE,
            "faktor":  cfg.FAKTOR_GORNJEG_PRAGA,
            "prag_d":  prag_d,
            "prag_g":  prag_g,
            "vad_db":  cfg.VAD_TOP_DB,
            "vad_min": cfg.VAD_MIN_DULJINA,
            "vad_sp":  cfg.VAD_SPAJANJE,
        }

        rez = pokreni_evaluaciju(gt, baza, prag_d, prag_g, params, tiho=False)
        ispisi_rezultat(rez)

        if args.xlsx:
            excel_put = os.path.join("rezultati", f"evaluacija_{ts}.xlsx")
            spremi_eval_excel(rez, excel_put)

    # ============================================================
    # MOD 2: SWEEP
    # ============================================================
    elif args.sweep:
        try:
            prop_vrijednosti   = [float(x) for x in args.prop_values.split(",")]
            faktor_vrijednosti = [float(x) for x in args.faktor_values.split(",")]
        except ValueError:
            print("  GREŠKA: Neispravne vrijednosti za sweep. Koristim defaulte.")
            prop_vrijednosti   = [0.0, 0.2, 0.4, 0.6, 0.75]
            faktor_vrijednosti = [1.5, 2.0, 2.5, 3.0]

        def _parse(s, tip=float):
            return [tip(x) for x in s.split(",")] if s else None

        rezultati = pokreni_sweep(
            gt, baza,
            prop_vrijednosti   = [float(x) for x in args.prop_values.split(",")],
            faktor_vrijednosti = [float(x) for x in args.faktor_values.split(",")],
            trajanje_vrijednosti    = _parse(args.trajanje_values),
            preklapanje_vrijednosti = _parse(args.prekl_values),
            vad_db_vrijednosti      = _parse(args.vad_db_values),
            vad_min_vrijednosti     = _parse(args.vad_min_values),
            vad_sp_vrijednosti      = _parse(args.vad_sp_values),
        )

        if args.xlsx:
            excel_put = os.path.join("rezultati", f"sweep_{ts}.xlsx")
            spremi_sweep_excel(rezultati, excel_put)

            # Spremi i detaljni prikaz najboljeg rezultata
            best_put = os.path.join("rezultati", f"sweep_best_{ts}.xlsx")
            spremi_eval_excel(rezultati[0], best_put)


if __name__ == "__main__":
    main()